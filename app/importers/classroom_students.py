from collections.abc import Sequence
from dataclasses import dataclass
from io import BytesIO
from typing import Never
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.cell import Cell, MergedCell
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from app.schemas.classroom_import import ClassroomStudentImportRow

MAX_XLSX_SIZE_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 1000
REQUIRED_HEADERS = ("full_name", "email")


@dataclass(frozen=True)
class ClassroomStudentImportIssue:
    row_number: int | None
    column: str | None
    message: str


class ClassroomStudentImportValidationError(Exception):
    def __init__(
        self,
        issues: Sequence[ClassroomStudentImportIssue],
    ) -> None:
        self.issues = tuple(issues)
        super().__init__("Classroom student import validation failed")


def parse_classroom_students_xlsx(
    content: bytes,
) -> list[ClassroomStudentImportRow]:
    try:
        return _parse_classroom_students_xlsx(content)
    except (
        BadZipFile,
        InvalidFileException,
        OSError,
        ValueError,
        KeyError,
        SyntaxError,
    ) as error:
        raise ClassroomStudentImportValidationError(
            [
                ClassroomStudentImportIssue(
                    row_number=None,
                    column=None,
                    message="The uploaded file is not a valid .xlsx workbook",
                )
            ]
        ) from error


def _parse_classroom_students_xlsx(
    content: bytes,
) -> list[ClassroomStudentImportRow]:
    if not content:
        _raise_single_issue("The Excel file is empty")

    if len(content) > MAX_XLSX_SIZE_BYTES:
        _raise_single_issue("The Excel file is larger than 5 MB")

    try:
        workbook = load_workbook(
            filename=BytesIO(content),
            read_only=True,
            data_only=False,
        )
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise ClassroomStudentImportValidationError(
            [
                ClassroomStudentImportIssue(
                    row_number=None,
                    column=None,
                    message="The uploaded file is not a valid .xlsx workbook",
                )
            ]
        ) from error

    try:
        if not workbook.worksheets:
            _raise_single_issue("The Excel workbook has no worksheets")

        worksheet = workbook.worksheets[0]
        rows = worksheet.iter_rows()
        header_row = next(rows, None)

        if header_row is None:
            _raise_single_issue("The Excel worksheet is empty")

        header_positions = _get_header_positions(header_row)
        issues: list[ClassroomStudentImportIssue] = []
        parsed_rows: list[ClassroomStudentImportRow] = []
        first_row_by_email: dict[str, int] = {}
        data_row_count = 0

        for row_number, row in enumerate(rows, start=2):
            if _row_is_empty(row):
                continue

            data_row_count += 1

            if data_row_count > MAX_IMPORT_ROWS:
                issues.append(
                    ClassroomStudentImportIssue(
                        row_number=row_number,
                        column=None,
                        message=f"An import can contain at most {MAX_IMPORT_ROWS} rows",
                    )
                )
                break

            row_data: dict[str, object] = {"row_number": row_number}
            row_has_formula = False

            for header in REQUIRED_HEADERS:
                cell = _get_cell(row, header_positions[header])

                if cell is not None and cell.data_type == "f":
                    issues.append(
                        ClassroomStudentImportIssue(
                            row_number=row_number,
                            column=header,
                            message="Formulas are not allowed",
                        )
                    )
                    row_has_formula = True
                    continue

                row_data[header] = None if cell is None else cell.value

            if row_has_formula:
                continue

            try:
                parsed_row = ClassroomStudentImportRow.model_validate(row_data)
            except ValidationError as error:
                issues.extend(_pydantic_issues(error, row_number))
                continue

            normalized_email = str(parsed_row.email)
            first_row_number = first_row_by_email.get(normalized_email)

            if first_row_number is not None:
                issues.append(
                    ClassroomStudentImportIssue(
                        row_number=row_number,
                        column="email",
                        message=f"Email duplicates row {first_row_number}",
                    )
                )
                continue

            first_row_by_email[normalized_email] = row_number
            parsed_rows.append(parsed_row)

        if data_row_count == 0:
            issues.append(
                ClassroomStudentImportIssue(
                    row_number=None,
                    column=None,
                    message="The Excel worksheet contains no student rows",
                )
            )

        if issues:
            raise ClassroomStudentImportValidationError(issues)

        return parsed_rows
    finally:
        workbook.close()


def _get_header_positions(
    header_row: Sequence[Cell | MergedCell],
) -> dict[str, int]:
    positions: dict[str, int] = {}
    duplicate_headers: set[str] = set()

    for index, cell in enumerate(header_row):
        header = _normalize_header(cell.value)

        if not header:
            continue

        if header in positions:
            duplicate_headers.add(header)
            continue

        positions[header] = index

    issues = [
        ClassroomStudentImportIssue(
            row_number=1,
            column=header,
            message="Header is duplicated",
        )
        for header in sorted(duplicate_headers)
    ]

    issues.extend(
        ClassroomStudentImportIssue(
            row_number=1,
            column=header,
            message="Required header is missing",
        )
        for header in REQUIRED_HEADERS
        if header not in positions
    )

    if issues:
        raise ClassroomStudentImportValidationError(issues)

    return positions


def _normalize_header(value: object) -> str:
    if not isinstance(value, str):
        return ""

    return value.strip().lower()


def _row_is_empty(row: Sequence[Cell | MergedCell]) -> bool:
    return all(
        cell.value is None or (isinstance(cell.value, str) and not cell.value.strip())
        for cell in row
    )


def _get_cell(
    row: Sequence[Cell | MergedCell],
    index: int,
) -> Cell | MergedCell | None:
    if index >= len(row):
        return None

    return row[index]


def _pydantic_issues(
    error: ValidationError,
    row_number: int,
) -> list[ClassroomStudentImportIssue]:
    issues: list[ClassroomStudentImportIssue] = []

    for item in error.errors(include_url=False):
        location = item["loc"]
        column = str(location[0]) if location else None

        issues.append(
            ClassroomStudentImportIssue(
                row_number=row_number,
                column=column,
                message=str(item["msg"]),
            )
        )

    return issues


def _raise_single_issue(message: str) -> Never:
    raise ClassroomStudentImportValidationError(
        [
            ClassroomStudentImportIssue(
                row_number=None,
                column=None,
                message=message,
            )
        ]
    )
